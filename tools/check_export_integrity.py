"""Independent file-open checks. Does not claim native Office validation."""
from pathlib import Path, PurePosixPath
import io
import posixpath
import zipfile
import xml.etree.ElementTree as ET
import fitz
from openpyxl import load_workbook
try:
    from pptx import Presentation
except ImportError:
    Presentation = None

ROOT = Path(__file__).resolve().parents[1]
ns = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main', 'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}

def check(path):
    data = path.read_bytes()
    if path.suffix == '.pdf':
        assert data.startswith(b'%PDF-') and data.rstrip().endswith(b'%%EOF'), path
        doc = fitz.open(stream=data, filetype='pdf')
        assert not doc.is_repaired and not doc.needs_pass and len(doc), path
        for page in doc:
            page.get_text()
            page.get_pixmap(matrix=fitz.Matrix(.25,.25))
        return f'{len(doc)} pages opened/rendered without repair'
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert z.testzip() is None, path
        names=set(z.namelist())
        assert '[Content_Types].xml' in names and '_rels/.rels' in names
        for name in names:
            if name.endswith(('.xml','.rels')):
                tree=ET.fromstring(z.read(name))
                if name.endswith('.rels'):
                    base='' if name=='_rels/.rels' else str(PurePosixPath(name).parent.parent)
                    ids=[]
                    for rel in tree:
                        ids.append(rel.attrib['Id'])
                        if rel.attrib.get('TargetMode')!='External':
                            target=posixpath.normpath(posixpath.join(base,rel.attrib['Target'])).lstrip('/')
                            assert target in names, (name,target)
                    assert len(ids)==len(set(ids)), name
        if path.suffix=='.pptx':
            master=ET.fromstring(z.read('ppt/slideMasters/slideMaster1.xml'))
            valid={'dk1','dk2','lt1','lt2','hlink','folHlink',*(f'accent{i}' for i in range(1,7))}
            assert all(v in valid for v in master.find('p:clrMap',ns).attrib.values())
            assert int(master.find('p:sldLayoutIdLst/p:sldLayoutId',ns).attrib['id'])>=2147483648
            theme=ET.fromstring(z.read('ppt/theme/theme1.xml'))
            for tag in ('fillStyleLst','lnStyleLst','effectStyleLst','bgFillStyleLst'):
                assert len(theme.find('.//a:'+tag,ns))==3,tag
    if path.suffix=='.xlsx':
        wb=load_workbook(io.BytesIO(data),data_only=False)
        assert len(wb.worksheets)>0
        return f'{len(wb.worksheets)} sheets opened'
    if Presentation is None:
        return 'ZIP/XML/relationships and OOXML-specific checks passed; desktop opening not tested'
    deck=Presentation(io.BytesIO(data))
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert isinstance(shape.text,str)
    return f'{len(deck.slides)} slides opened'

if __name__=='__main__':
    files=sorted((ROOT/'.export-validation').glob('*'))+sorted((ROOT/'output/pdf').glob('*.pdf'))
    for path in files:
        if path.suffix in ('.xlsx','.pptx','.pdf'):
            print('PASS',path.name,check(path))
